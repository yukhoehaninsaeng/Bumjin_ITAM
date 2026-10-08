"""Run only with the web app stopped. Explicit, read-only directory import."""
import os, ssl
import ldap3
from app import FILE, LOCK, digest, save, literal
import openpyxl

if __name__=='__main__':
    if FILE.with_suffix('.app.lock').exists():raise SystemExit('웹앱을 먼저 종료하세요.')
    required=['LDAP_BIND_DN','LDAP_BIND_PASSWORD','LDAP_NAME_ATTRIBUTE','LDAP_ID_ATTRIBUTE','LDAP_USER_FILTER']
    for key in required:
        if not os.environ.get(key):raise SystemExit(key+' 환경변수를 설정하세요. 속성명/필터는 실제 조회 후 확정해야 합니다.')
    host=os.environ.get('LDAP_HOST','ldap.bumjin.local')
    tls=ldap3.Tls(validate=ssl.CERT_REQUIRED,ca_certs_file=os.environ.get('LDAP_CA_FILE') or None)
    server=ldap3.Server(host,port=636,use_ssl=True,tls=tls,connect_timeout=10)
    conn=ldap3.Connection(server,user=os.environ['LDAP_BIND_DN'],password=os.environ['LDAP_BIND_PASSWORD'],auto_bind=True,receive_timeout=20,raise_exceptions=True)
    attrs=['uid','mail',os.environ['LDAP_NAME_ATTRIBUTE'],os.environ['LDAP_ID_ATTRIBUTE']]
    entries=list(conn.extend.standard.paged_search(os.environ.get('LDAP_SEARCH_BASE','cn=users,dc=ldap,dc=bumjin,dc=local'),os.environ['LDAP_USER_FILTER'],attributes=attrs,paged_size=100,generator=True))
    if conn.result.get('result')!=0:raise SystemExit('LDAP 검색 실패. 기존 목록 유지.')
    def first(v):return str(v[0] if isinstance(v,list) and v else v or '')
    users=[]
    for entry in entries:
        if entry['type']!='searchResEntry':continue
        a=entry['attributes']; uid=first(a.get('uid')); stable=first(a.get(os.environ['LDAP_ID_ATTRIBUTE']))
        if uid and stable:users.append([uid,first(a.get(os.environ['LDAP_NAME_ATTRIBUTE'])),first(a.get('mail')),stable])
        elif uid:raise SystemExit('안정적 ID가 없는 계정 발견. 기존 목록 유지.')
    if not users:raise SystemExit('조회된 사용자가 없습니다. 기존 목록 유지.')
    old=digest();w=openpyxl.load_workbook(FILE)
    s=w['앱_LDAP사용자']
    previous={str(r[0]):str(r[3]) for r in s.iter_rows(min_row=2,max_col=4,values_only=True) if r[0]}
    if any(uid in previous and previous[uid]!=stable for uid,_,_,stable in users):raise SystemExit('계정 ID 재사용 후보 발견. 자동 매핑 금지, 검토 필요.')
    print('조회된 계정:',len(users),'기존 계정:',len(previous))
    print('필터가 직원/활성 계정만 포함하는지 검증해야 합니다. 이번 버전은 사용자 선택 목록을 가져오며 자산 배정을 자동 변경하지 않습니다.')
    if input('적용하려면 APPLY 입력: ')!='APPLY':raise SystemExit('취소')
    s.delete_rows(2,s.max_row)
    for row in users:
        r=s.max_row+1
        for c,v in enumerate(row,1):literal(s.cell(r,c),v)
    save(w,old);conn.unbind();print('직원 선택 목록을 엑셀에 저장했습니다. 웹앱을 다시 실행하세요.')
