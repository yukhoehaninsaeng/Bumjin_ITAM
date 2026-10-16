import io
import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import MagicMock, patch
import openpyxl
import app
import auth_service
import admin_service

ASSET='11111111-1111-4111-8111-111111111111'

class ReadOnlyTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.original=app.FILE;app.FILE=Path(self.temp.name)/'test.xlsx'
        self.environ=patch.dict(os.environ,{'APP_PASSWORD':'test-admin-password','PUBLIC_BASE_URL':'http://example.invalid','BACKEND_GATEWAY_TOKEN':''});self.environ.start()
        auth_service.SESSIONS.clear();auth_service.PREAUTH.clear();auth_service.ATTEMPTS.clear()
        w=openpyxl.Workbook()
        for name,*_ in app.CONFIG.values():w.create_sheet(name)
        s=w[app.CONFIG['BJE'][0]]
        for col,val in {'B':'TEST-D-001','C':'구매팀','D':'test-account','G':'Test Laptop','H':'PRIVATE-SERIAL','Q':'PRIVATE-SW','X':'PRIVATE-NOTE','Y':'2026-01-01','Z':'사용','AZ':ASSET,'BA':'test-account'}.items():s[f'{col}6']=val
        w.create_sheet('앱_LDAP사용자').append(['uid','name','email','id','department','dn','description'])
        w['앱_LDAP사용자'].append(['test-account','test-account','PRIVATE-EMAIL','PRIVATE-UUID','BJE/구매팀','PRIVATE-DN','테스트 직원'])
        w.create_sheet('앱_이력').append(['시각','자산ID','유형','이전','이후','작업자']);w.save(app.FILE);w.close()
    def tearDown(self):
        app.FILE=self.original;self.environ.stop();self.temp.cleanup()
    def request(self,path,method='GET',payload=None,cookie='',csrf='',query=''):
        body=json.dumps(payload).encode() if payload is not None else b'';captured={}
        env={'PATH_INFO':path,'QUERY_STRING':query,'REQUEST_METHOD':method,'HTTP_COOKIE':cookie,'HTTP_X_CSRF_TOKEN':csrf,'CONTENT_LENGTH':str(len(body)),'wsgi.input':io.BytesIO(body),'REMOTE_ADDR':'test','wsgi.url_scheme':'http'}
        def start(code,headers):captured.update(code=code,headers=dict(headers))
        result=b''.join(app.application(env,start));return captured['code'],captured['headers'],result
    def session(self,role):
        token,record=auth_service.issue_session(app,'admin' if role=='admin' else 'test-account',role)
        return auth_service.COOKIE+'='+token,record['csrf']
    def test_direct_qr_is_readonly_even_for_admin(self):
        for cookie in ('',self.session('admin')[0],self.session('viewer')[0]):
            code,_,body=self.request('/a/'+ASSET,cookie=cookie)
            self.assertEqual(code,'200 OK');self.assertIn('조회 전용'.encode(),body);self.assertNotIn(b'id="collectAsset"',body)
        self.assertEqual(self.request('/a/not-an-asset')[0],'404 Not Found')
    def test_viewer_can_open_admin_login_without_being_redirected_to_forbidden_page(self):
        cookie,_=self.session('viewer')
        code,_,body=self.request('/login',cookie=cookie,query='next=%2Fmanage%2Fa%2F'+ASSET)
        self.assertEqual(code,'200 OK');self.assertIn(b'loginForm',body)

    def test_public_asset_allowlist(self):
        code,_,body=self.request('/api/public/assets/'+ASSET)
        self.assertEqual(code,'200 OK');asset=json.loads(body)['asset']
        self.assertEqual(asset['user'],'테스트 직원');self.assertEqual(asset['department'],'BJE/구매팀')
        for key in ('uid','email','serial','note','office','dn','history','revision','csrf','directory_id'):self.assertNotIn(key,asset)
        for secret in (b'PRIVATE',b'test-account'):self.assertNotIn(secret,body)
        self.assertEqual(self.request('/api/public/assets/22222222-2222-4222-8222-222222222222')[0],'404 Not Found')
    def test_viewer_cannot_read_admin_data_or_write_with_valid_csrf(self):
        cookie,csrf=self.session('viewer')
        for path in ('/api/data','/api/admin','/api/maintenance','/api/export','/qr/'+ASSET,'/users','/admin','/manage/a/'+ASSET):
            self.assertEqual(self.request(path,cookie=cookie)[0],'403 Forbidden',path)
        for path in ('/api/assets','/api/assets/map-user','/api/organizations','/api/admin','/api/admin/ldap','/api/maintenance','/api/labels/pdf'):
            self.assertEqual(self.request(path,'POST',{'role':'admin'},cookie,csrf)[0],'403 Forbidden',path)
        self.assertEqual(self.request('/api/view/assets',cookie=cookie)[0],'200 OK')
        self.assertEqual(self.request('/',cookie=cookie)[0],'200 OK')
        self.assertEqual(self.request('/api/auth/logout','POST',{},cookie,csrf)[0],'200 OK')
    def test_guest_cannot_write_or_list(self):
        self.assertEqual(self.request('/api/assets','POST',{})[0],'401 Unauthorized')
        self.assertEqual(self.request('/api/view/assets')[0],'401 Unauthorized')
        code,headers,_=self.request('/manage/a/'+ASSET)
        self.assertEqual(code,'303 See Other');self.assertIn('/login?next=',headers['Location'])
    def test_admin_manage_and_actor_cannot_be_spoofed(self):
        cookie,csrf=self.session('admin')
        self.assertEqual(self.request('/manage/a/'+ASSET,cookie=cookie)[0],'200 OK')
        with patch.object(app,'mutate',return_value={'id':ASSET}) as mutate:
            self.assertEqual(self.request('/api/assets','POST',{'_actor':'forged'},cookie,csrf)[0],'200 OK')
            self.assertEqual(mutate.call_args.args[0]['_actor'],'admin')
        self.assertEqual(self.request('/api/assets','POST',{},cookie)[0],'403 Forbidden')
    def test_ldap_login_is_viewer_and_returns_to_readonly_asset(self):
        body,header=auth_service.issue_preauth({});cookie=header[1].split(';')[0]
        env={'HTTP_COOKIE':cookie,'HTTP_X_CSRF_TOKEN':body['csrf'],'REMOTE_ADDR':'test'}
        with patch.object(auth_service,'ldap_password_valid',return_value=True):
            code,result,headers=auth_service.login(app,env,{'username':'employee','password':'secret','role':'admin','next':'/manage/a/'+ASSET})
        self.assertEqual(code,'200 OK');self.assertEqual(result['role'],'viewer');self.assertEqual(result['next'],'/a/'+ASSET)
        session_cookie=next(v.split(';')[0] for k,v in headers if v.startswith(auth_service.COOKIE+'='))
        self.assertEqual(auth_service.current(app,{'HTTP_COOKIE':session_cookie})['role'],'viewer')
    def test_ldap_login_binds_selected_live_dn_and_escapes_search(self):
        cfg={'host':'test','port':636,'base_dn':'dc=test','user_base':'','user_filter':'(objectClass=posixAccount)','uid_attr':'uid','id_attr':'entryUUID'}
        service=MagicMock();service.result={'result':0};service.response=[{'type':'searchResEntry','dn':'uid=employee,dc=test','attributes':{'uid':['employee'], 'entryUUID':['stable']}}]
        user_connection=MagicMock();user_connection.bound=True
        with patch.object(admin_service,'settings',return_value={'ldap':cfg}),patch.object(admin_service,'connect_directory',return_value=service),patch('ldap3.Connection',return_value=user_connection) as bind:
            self.assertTrue(auth_service.ldap_password_valid(app,'employee','secret'))
            self.assertEqual(bind.call_args.kwargs['user'],'uid=employee,dc=test');self.assertEqual(bind.call_args.kwargs['password'],'secret')
            service.response=[];self.assertFalse(auth_service.ldap_password_valid(app,'*)(uid=*)','secret'))
            self.assertIn('\\2a',service.search.call_args.args[1])
        service.unbind.assert_called();user_connection.unbind.assert_called()

if __name__=='__main__':unittest.main()
