"""Bumjin Excel-backed pilot. Single process, no database server."""
import signal, re, base64, copy, datetime as dt, hashlib, hmac, io, json, os, secrets, shutil, threading, uuid
from pathlib import Path
from urllib.parse import urlsplit, quote, parse_qs
from wsgiref.simple_server import make_server
import openpyxl
import admin_service
import auth_service
import organization_service
import label_pdf
from openpyxl.workbook.properties import CalcProperties

ROOT = Path(__file__).resolve().parent
FILE = Path(os.environ.get('ASSET_FILE', str(ROOT/'data/Bumjin_IT자산관리대장.xlsx')))
LOCK = threading.RLock()
CONFIG = {'BJM': ('3.BJM PC 현황','W','X','U','V'), 'BJE': ('4.BJE PC 현황','Y','Z','V','X'), 'BJC': ('5.BJC PC 현황','Y','Z','V','X')}
CONFIG.update({'BJM:모니터':('3_1.BJM 모니터 현황','I','K','H','J'),'BJE:모니터':('4_1.BJE 모니터 현황','I','K','H','J'),'BJC:모니터':('5_1.BJC 모니터 현황','I','K','H','J')})
MONITOR_BASE={'number':'B','department':'C','user':'D','size':'E','model':'F','serial':'G'}
BASE = {'number':'B','department':'C','user':'D','type':'E','maker':'F','model':'G','serial':'H','cpu':'I','bits':'J','ram':'K','ssd':'L','hdd':'M','gpu':'N','os':'O','edition':'P','office':'Q','hangul':'R'}

SOFTWARE_EXTRA={'BJM':{'officekeeper':'S','antivirus':'T'},'BJE':{'sap':'S','antivirus':'T','docuone':'U','other':'W'},'BJC':{'sap':'S','antivirus':'T','docuone':'U','other':'W'}}

def digest(): return hashlib.sha256(FILE.read_bytes()).hexdigest()
def value(v): return v.isoformat()[:10] if isinstance(v,(dt.datetime,dt.date)) else str(v or '')
def literal(cell, v):
    cell.value = str(v or '')[:2000]
    cell.data_type = 's'  # user data cannot become an Excel formula
def save(w, expected):
    if digest()!=expected: raise ValueError('엑셀이 외부에서 변경됐습니다. 새로고침 후 다시 시도하세요.')
    directory=FILE.parent/'backups'; directory.mkdir(exist_ok=True)
    shutil.copy2(FILE,directory/(dt.datetime.now().strftime('%Y%m%d_%H%M%S_%f')+'.xlsx'))
    temp=FILE.with_name(FILE.stem+'.pending.xlsx')
    try:
        w.calculation=CalcProperties(fullCalcOnLoad=True)
        w.save(temp)
        check=openpyxl.load_workbook(temp); check.close()
        if digest()!=expected: raise ValueError('저장 중 외부 수정 감지. 앱 쓰기를 중지하고 파일을 확인하세요.')
        os.replace(temp,FILE)
    finally: temp.unlink(missing_ok=True)

def initialize():
    with LOCK:
        old=digest(); w=openpyxl.load_workbook(FILE); changed=False
        for division,(name,*_) in CONFIG.items():
            s=w[name]
            if s['AZ5'].value not in (None,'APP_ASSET_ID'): raise ValueError('AZ열이 이미 사용 중입니다.')
            if s['BA5'].value not in (None,'APP_LDAP_UID'): raise ValueError('BA열이 이미 사용 중입니다.')
            for col,label in [('AZ','APP_ASSET_ID'),('BA','APP_LDAP_UID')]:
                if s[f'{col}5'].value!=label: s[f'{col}5']=label; changed=True
                s.column_dimensions[col].hidden=True
            for r in range(6,s.max_row+1):
                if s[f'B{r}'].value and not s[f'AZ{r}'].value:
                    s[f'AZ{r}']=str(uuid.uuid4()); changed=True
        for name,headers in [('앱_이력',['시각','자산ID','유형','이전','이후']),('앱_LDAP사용자',['uid','실명','이메일','디렉터리ID'])]:
            if name not in w: w.create_sheet(name).append(headers); changed=True
        if organization_service.bootstrap_from_excel(__import__(__name__),w):changed=True
        if changed: save(w,old)

def assets(w):
    out=[]
    for division,(name,date,status,grade,note) in CONFIG.items():
        s=w[name]
        for r in range(6,s.max_row+1):
            if not s[f'B{r}'].value: continue
            monitor=':모니터' in division
            a={k:'' for k in BASE}
            a.update({k:value(s[f'{c}{r}'].value) for k,c in (MONITOR_BASE if monitor else BASE).items()})
            a.update(id=value(s[f'AZ{r}'].value),uid=value(s[f'BA{r}'].value),division=division.split(':')[0],category='모니터' if monitor else 'PC',row=r,
                     date=value(s[f'{date}{r}'].value),status=value(s[f'{status}{r}'].value),note=value(s[f'{note}{r}'].value),
                     date_label='제조일자' if monitor else {'BJM':'구매일자','BJE':'구입년월','BJC':'제조일자'}[division])
            if not monitor: a.update({k:value(s[f'{c}{r}'].value) for k,c in SOFTWARE_EXTRA[division].items()})
            a['grade']='미확인'
            try:
                d=dt.date.fromisoformat(a['date']); t=dt.date.today(); years=t.year-d.year-((t.month,t.day)<(d.month,d.day))
                if years>=0:a['grade']='D' if years>=10 else 'C' if years>=5 else 'B' if years>=3 else 'A'
            except ValueError: pass
            out.append(a)
    return out

def snapshot():
    with LOCK:
        rev=digest(); w=openpyxl.load_workbook(FILE)
        users=[dict(zip(['uid','name','email','directory_id','department','dn','description'],map(value,r))) for r in w['앱_LDAP사용자'].iter_rows(min_row=2,max_col=7,values_only=True) if r[0]]
        history=[list(map(value,r)) for r in w['앱_이력'].iter_rows(min_row=2,max_col=5,values_only=True)]
        return {'assets':assets(w),'users':users,'revision':rev,'history':history[-300:][::-1],
                'organizations':[{'name':value(r[0]),'dn':value(r[1])} for r in w['앱_LDAP조직'].iter_rows(min_row=2,max_col=2,values_only=True) if r[0]] if '앱_LDAP조직' in w else [],'managed_organizations':organization_service.rows(w),'asset_organizations':organization_service.asset_links(w),'organization_history':organization_service.history(w),'schema_version':2,'base_url':os.environ.get('PUBLIC_BASE_URL','http://127.0.0.1:1880')}


def next_number(all_assets, division, category, device_type='Desktop'):
    if division not in ('BJM','BJE','BJC') or category not in ('PC','모니터'): raise ValueError('법인과 장비 구분을 선택하세요.')
    kind='M' if category=='모니터' else ('N' if any(x in str(device_type).lower() for x in ['notebook','laptop','노트북']) else 'D')
    prefix=f'{division}-{kind}-'
    maximum=max((int(a['number'][len(prefix):]) for a in all_assets if a['number'].startswith(prefix) and a['number'][len(prefix):].isdigit()),default=0)
    return prefix+str(maximum+1).zfill(3)

def mutate(payload):
    with LOCK:
        old=digest()
        if payload.get('revision')!=old: raise ValueError('다른 변경이 있습니다. 새로고침 후 저장하세요.')
        w=openpyxl.load_workbook(FILE); all_assets=assets(w); aid=payload.get('id'); before={}
        if aid:
            before=next((a for a in all_assets if a['id']==aid),None)
            if not before: raise ValueError('자산이 없습니다.')
            division=before['division']+(':모니터' if before.get('category')=='모니터' else ''); row=before['row']
            action=payload.get('action','')
            if action not in ('','수거','사용자 변경','사용자 확인'): raise ValueError('작업 유형이 올바르지 않습니다.')
            payload={**before,**payload}
            if action=='수거': payload.update(user='',uid='',status='보관')
            if action=='사용자 변경':
                if not str(payload.get('user','')).strip() and not str(payload.get('uid','')).strip(): raise ValueError('새 사용자를 입력하세요.')
                payload['status']='사용'
            if action=='사용자 확인' and not before['user']: raise ValueError('현재 배정된 사용자가 없습니다.')
        else:
            division=payload.get('division','')+(':모니터' if payload.get('category')=='모니터' else '')
            if division not in CONFIG: raise ValueError('사업부를 선택하세요.')
            used=[a['row'] for a in all_assets if a['division']==division.split(':')[0] and a['category']==('모니터' if ':모니터' in division else 'PC')]; row=max(used,default=5)+1; aid=str(uuid.uuid4())
        name,date,status,grade,note=CONFIG[division]; s=w[name]
        number=next_number(all_assets,division.split(':')[0],'모니터' if ':모니터' in division else 'PC',payload.get('type','Desktop')) if not before and payload.get('auto_number') else str(payload.get('number','')).strip()
        if not number: raise ValueError('관리번호는 필수입니다.')
        payload['number']=number
        if (not before or number!=before['number']) and any(a['number']==number and a['id']!=aid for a in all_assets): raise ValueError('동일 관리번호가 있습니다.')
        state=payload.get('status','보관')
        if state not in ['사용','보관','수리','폐기','분실']: raise ValueError('상태 값이 올바르지 않습니다.')
        datevalue=str(payload.get('date','')).strip()
        parsed=dt.date.fromisoformat(datevalue) if datevalue else None
        uid=str(payload.get('uid','')).strip(); username=str(payload.get('user','')).strip()
        if uid:
            matches=[r for r in w['앱_LDAP사용자'].iter_rows(min_row=2,values_only=True) if value(r[0])==uid]
            if len(matches)!=1: raise ValueError('LDAP 계정 선택이 올바르지 않습니다.')
            username=(value(matches[0][6]) if len(matches[0])>6 else '') or value(matches[0][1]) or uid
        if not before:
            for c in range(1,28):
                src=s.cell(row-1,c); dst=s.cell(row,c)
                if src.has_style: dst._style=copy.copy(src._style)
        for k,c in (MONITOR_BASE if ':모니터' in division else BASE).items(): literal(s[f'{c}{row}'],username if k=='user' else payload.get(k,''))
        if ':모니터' not in division:
            for k,c in SOFTWARE_EXTRA[division].items():
                if k in payload: literal(s[f'{c}{row}'],payload[k])
        literal(s[f'{note}{row}'],payload.get('note',''));literal(s[f'{status}{row}'],state)
        s[f'{date}{row}']=parsed; s[f'{date}{row}'].number_format='yyyy-mm-dd'
        s[f'{grade}{row}']=f'=IF({date}{row}="","",IF(DATEDIF({date}{row},TODAY(),"Y")>=10,"D",IF(DATEDIF({date}{row},TODAY(),"Y")>=5,"C",IF(DATEDIF({date}{row},TODAY(),"Y")>=3,"B","A"))))'
        s[f'AZ{row}']=aid;literal(s[f'BA{row}'],uid)
        after=next(a for a in assets(w) if a['id']==aid)
        w['앱_이력'].append([dt.datetime.now().astimezone().isoformat(),aid,payload.get('action') or ('수정' if before else '등록'),json.dumps(before,ensure_ascii=False),json.dumps(after,ensure_ascii=False)])
        save(w,old)
        return {'id':aid}


def map_user_assets(payload):
    """Bind selected name-only assets to one imported LDAP identity in one save."""
    with LOCK:
        old=digest()
        if payload.get('revision')!=old:raise ValueError('다른 변경이 있습니다. 새로고침 후 매핑하세요.')
        ids=payload.get('asset_ids');source=str(payload.get('source_name',''))
        if not isinstance(ids,list) or not 1<=len(ids)<=500 or any(not isinstance(i,str) for i in ids) or len(set(ids))!=len(ids):
            raise ValueError('매핑할 자산 1~500개를 중복 없이 선택하세요.')
        if not source.strip():raise ValueError('기존 사용자 이름이 없습니다.')
        w=openpyxl.load_workbook(FILE)
        target=str(payload.get('target_directory_id',''))
        users=[list(r) for r in w['앱_LDAP사용자'].iter_rows(min_row=2,max_col=7,values_only=True) if target and value(r[3])==target]
        if len(users)!=1 or not users[0][0]:raise ValueError('LDAP 계정을 다시 선택하세요. 사용자 목록을 동기화해야 할 수 있습니다.')
        user=users[0];uid=value(user[0]);name=value(user[6]) or value(user[1]) or uid
        if sum(value(r[0])==uid for r in w['앱_LDAP사용자'].iter_rows(min_row=2,max_col=1,values_only=True))!=1:
            raise ValueError('LDAP 계정이 중복됩니다. 다시 동기화하세요.')
        index={a['id']:a for a in assets(w)}
        if any(i not in index for i in ids):raise ValueError('선택한 자산이 없습니다. 새로고침하세요.')
        selected=[index[i] for i in ids]
        if any(a['uid'] or a['user']!=source for a in selected):
            raise ValueError('이미 매핑됐거나 사용자 이름이 다른 자산이 포함돼 있습니다. 새로고침 후 다시 선택하세요.')
        timestamp=dt.datetime.now().astimezone().isoformat()
        for before in selected:
            division=before['division']+(':모니터' if before['category']=='모니터' else '')
            sheet=w[CONFIG[division][0]];row=before['row']
            literal(sheet[f'BA{row}'],uid);literal(sheet[f'D{row}'],name)
            after={**before,'uid':uid,'user':name}
            event=[timestamp,before['id'],'LDAP 일괄 매핑',json.dumps(before,ensure_ascii=False),json.dumps(after,ensure_ascii=False)]
            r=w['앱_이력'].max_row+1
            for c,v in enumerate(event,1):literal(w['앱_이력'].cell(r,c),v)
        save(w,old)
        return {'mapped':len(selected),'uid':uid,'directory_id':target,'name':name}


MAINT_FIELDS={'vendor':'B','service':'C','quantity':'D','division':'E','period':'F','cost2024':'G','cost2025':'H','cost2026':'I','contact':'J','note':'L','end':'M','start':'N'}
def maintenance_snapshot():
    with LOCK:
        w=openpyxl.load_workbook(FILE);s=w['2.전산 유지보수']
        def resolved(row,col):
            cell=s.cell(row,col)
            for area in s.merged_cells.ranges:
                if area.min_row<=row<=area.max_row and area.min_col<=col<=area.max_col:
                    return s.cell(area.min_row,area.min_col).value
            return cell.value
        contracts=[]
        for row in range(3,18):
            if not s.cell(row,5).value: continue
            a={k:value(resolved(row,openpyxl.utils.column_index_from_string(c))) for k,c in MAINT_FIELDS.items()};a['row']=row
            try:a['remaining']=(dt.date.fromisoformat(a['end'])-dt.date.today()).days
            except ValueError:a['remaining']=None
            a['shared']=[k for k,c in MAINT_FIELDS.items() if isinstance(s[f'{c}{row}'],openpyxl.cell.cell.MergedCell)]
            contracts.append(a)
        if '앱_유지보수' in w:
            extra=w['앱_유지보수']
            for r in range(2,extra.max_row+1):
                a={k:value(extra[f'{c}{r}'].value) for k,c in MAINT_FIELDS.items()}
                if not a['service']:continue
                a.update(row=1000+r,shared=[])
                try:a['remaining']=(dt.date.fromisoformat(a['end'])-dt.date.today()).days
                except ValueError:a['remaining']=None
                contracts.append(a)
        sections=[];section=None
        for row in range(18,s.max_row+1):
            vals=[value(resolved(row,c)) for c in range(2,18)]
            title=next((v for v in vals if v.startswith('◎')),None)
            if title:section={'title':title,'rows':[]};sections.append(section)
            elif section and any(vals):section['rows'].append(vals)
        return {'contracts':contracts,'sections':sections,'revision':digest(),'csrf':SESSION}
def mutate_maintenance(payload):
    with LOCK:
        old=digest()
        if payload.get('revision')!=old:raise ValueError('다른 변경이 있습니다. 새로고침하세요.')
        w=openpyxl.load_workbook(FILE);s=w['2.전산 유지보수'];row=int(payload.get('row',0))
        if row==0 or row>=1002:
            if '앱_유지보수' not in w:
                w.create_sheet('앱_유지보수')
                for k,c in MAINT_FIELDS.items():w['앱_유지보수'][f'{c}1']=k
            s=w['앱_유지보수'];row=s.max_row+1 if row==0 else row-1000
            if row<2 or row>s.max_row+1:raise ValueError('계약 행이 없습니다.')
        elif row not in range(3,18):raise ValueError('원본 계약 행을 선택하세요.')
        if not str(payload.get('service',s[f'C{row}'].value or '')).strip():raise ValueError('서비스 이름은 필수입니다.')
        for k,c in MAINT_FIELDS.items():
            if k not in payload:continue
            cell=s[f'{c}{row}']
            if isinstance(cell,openpyxl.cell.cell.MergedCell):continue
            v=payload[k]
            if k in ('start','end'):cell.value=dt.date.fromisoformat(v) if v else None;cell.number_format='yyyy-mm-dd'
            elif k.startswith('cost'):
                n=float(v or 0)
                if n<0:raise ValueError('비용은 0 이상이어야 합니다.')
                cell.value=n
            else:literal(cell,v)
        s[f'K{row}']=f'=IF(M{row}="","",M{row}-TODAY())'
        w['앱_이력'].append([dt.datetime.now().astimezone().isoformat(),'maintenance:'+str(row),'유지보수 수정','',json.dumps(payload,ensure_ascii=False)])
        save(w,old);return {'row':row}

def qr_svg(url):
    from reportlab.graphics.barcode.qr import QrCodeWidget
    from reportlab.graphics.shapes import Drawing
    from reportlab.graphics import renderSVG
    qr=QrCodeWidget(url,barLevel='H',barBorder=4);x,y,x2,y2=qr.getBounds();d=Drawing(240,240,transform=[240/(x2-x),0,0,240/(y2-y),0,0]);d.add(qr)
    # Central logo covers a small horizontal strip; quiet zone is preserved.
    svg=renderSVG.drawToString(d).replace('<title>', '<rect x="0" y="0" width="240" height="240" fill="white"/><title>')
    logo_path=ROOT.parent/'public/qr-label-logo.svg'
    if logo_path.exists():
        logo=logo_path.read_text('utf-8')
        match=re.search(r'xlink:href="([^"]+)"',logo)
        if match:
            svg=svg.replace('</svg>','<rect x="49" y="102" width="142" height="36" rx="2" fill="white"/><image x="52" y="105" width="136" height="30" preserveAspectRatio="xMidYMid meet" xlink:href="'+match.group(1)+'"/></svg>')
    return svg.encode()

SESSION=secrets.token_urlsafe(32)
def application(env,start):
    def reply(code,data,ctype='application/json; charset=utf-8',extra=None):
        if not isinstance(data,bytes): data=json.dumps(data,ensure_ascii=False).encode()
        start(code,[('Content-Type',ctype),('Cache-Control','no-store'),('X-Content-Type-Options','nosniff'),('X-Frame-Options','DENY')]+(extra or []));return [data]
    gateway=os.environ.get('BACKEND_GATEWAY_TOKEN','')
    if gateway and not hmac.compare_digest(env.get('HTTP_X_GATEWAY_TOKEN',''),gateway):
        return reply('403 Forbidden',{'error':'게이트웨이 인증이 필요합니다.'})
    path=env.get('PATH_INFO','/');method=env.get('REQUEST_METHOD','GET');module=__import__(__name__)
    session=auth_service.current(module,env)
    if method=='GET' and path=='/login-logo.svg':
        return reply('200 OK',(ROOT.parent/'public/qr-label-logo.svg').read_bytes(),'image/svg+xml')
    if method=='GET' and path=='/api/auth/csrf':
        body,cookie=auth_service.issue_preauth(env);return reply('200 OK',body,extra=[cookie])
    if method=='POST' and path=='/api/auth/login':
        try:
            length=int(env.get('CONTENT_LENGTH') or '0')
            if not 0<length<=4096:return reply('400 Bad Request',{'error':'로그인 요청이 올바르지 않습니다.'})
            payload=json.loads(env['wsgi.input'].read(length))
            if not isinstance(payload,dict):raise ValueError()
        except (ValueError,KeyError):return reply('400 Bad Request',{'error':'로그인 요청이 올바르지 않습니다.'})
        code,body,headers=auth_service.login(module,env,payload);return reply(code,body,extra=headers)
    if method=='GET' and path=='/login':
        if session:
            target=auth_service.safe_next(parse_qs(env.get('QUERY_STRING','')).get('next',['/'])[0])
            return reply('303 See Other',b'',extra=[('Location',target)])
        return reply('200 OK',(ROOT.parent/'public/login.html').read_bytes(),'text/html; charset=utf-8')
    if not session:
        if method=='GET' and (path in ('/','/admin','/users') or path.startswith('/a/')):
            return reply('303 See Other',b'',extra=[('Location','/login?next='+quote(path,safe=''))])
        return reply('401 Unauthorized',{'error':'로그인이 필요합니다.'})
    request_csrf=session['csrf']
    if method=='POST' and path=='/api/auth/logout':
        if not hmac.compare_digest(env.get('HTTP_X_CSRF_TOKEN',''),request_csrf):return reply('403 Forbidden',{'error':'새로고침 후 다시 시도하세요.'})
        auth_service.remove_session(env)
        return reply('200 OK',{'ok':True},extra=[auth_service.cookie_header(env,auth_service.COOKIE,'',0)])
    try:
        if method=='POST' and path=='/api/labels/pdf':
            if not hmac.compare_digest(env.get('HTTP_X_CSRF_TOKEN',''),request_csrf):return reply('403 Forbidden',{'error':'새로고침하세요.'})
            length=int(env.get('CONTENT_LENGTH') or '0')
            if length>100000:return reply('413 Payload Too Large',{})
            payload=json.loads(env['wsgi.input'].read(length));ids=payload.get('ids',[])
            if not isinstance(ids,list) or not 1<=len(ids)<=500 or any(not isinstance(i,str) for i in ids) or len(set(ids))!=len(ids):raise ValueError('고유한 자산 1~500개를 선택하세요.')
            current={a['id']:a for a in snapshot()['assets']}
            if any(i not in current for i in ids):raise ValueError('선택한 자산이 없습니다. 새로고침하세요.')
            pdf=label_pdf.make_pdf([current[i] for i in ids],os.environ.get('PUBLIC_BASE_URL','http://127.0.0.1:1880'),ROOT.parent/'public/qr-label-logo.svg',payload)
            return reply('200 OK',pdf,'application/pdf',[('Content-Disposition','attachment; filename="Bumjin_QRP-3990.pdf"')])
        if method=='GET' and path=='/api/admin':return reply('200 OK',{**admin_service.settings(__import__(__name__)),'csrf':request_csrf})
        if method=='POST' and path in ('/api/admin','/api/admin/ldap','/api/organizations'):
            if not hmac.compare_digest(env.get('HTTP_X_CSRF_TOKEN',''),request_csrf):return reply('403 Forbidden',{'error':'새로고침하세요.'})
            length=int(env.get('CONTENT_LENGTH') or '0')
            if length>20000:return reply('413 Payload Too Large',{})
            payload=json.loads(env['wsgi.input'].read(length))
            try:
                module=__import__(__name__)
                result=organization_service.mutate(module,payload) if path=='/api/organizations' else admin_service.synchronize(module,payload) if path.endswith('/ldap') else admin_service.update_settings(module,payload)
            except ValueError:raise
            except Exception:return reply('400 Bad Request',{'error':'LDAP 연결 또는 설정 처리 실패. 주소·인증서·계정·검색 권한을 확인하세요.'})
            return reply('200 OK',result)
        if method=='GET' and path=='/api/maintenance':return reply('200 OK',{**maintenance_snapshot(),'csrf':request_csrf})
        if method=='POST' and path=='/api/maintenance':
            if not hmac.compare_digest(env.get('HTTP_X_CSRF_TOKEN',''),request_csrf):return reply('403 Forbidden',{'error':'새로고침하세요.'})
            length=int(env.get('CONTENT_LENGTH') or '0')
            if length>20000:return reply('413 Payload Too Large',{})
            return reply('200 OK',mutate_maintenance(json.loads(env['wsgi.input'].read(length))))
        if method=='GET' and path=='/api/data':return reply('200 OK',{**snapshot(),'csrf':request_csrf})
        if method=='GET' and path=='/api/export':return reply('200 OK',FILE.read_bytes(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', [('Content-Disposition','attachment; filename="Bumjin_IT_assets.xlsx"')])
        if method=='GET' and path.startswith('/qr/'):
            aid=path.split('/')[-1]
            if not any(a['id']==aid for a in snapshot()['assets']):return reply('404 Not Found',{})
            url=os.environ.get('PUBLIC_BASE_URL','http://127.0.0.1:1880').rstrip('/')+'/a/'+aid
            return reply('200 OK',qr_svg(url),'image/svg+xml')
        if method=='POST' and path in ('/api/assets','/api/assets/map-user'):
            if not hmac.compare_digest(env.get('HTTP_X_CSRF_TOKEN',''),request_csrf):return reply('403 Forbidden',{'error':'세션이 변경됐습니다. 새로고침하세요.'})
            length=int(env.get('CONTENT_LENGTH') or '0')
            if length>100000:return reply('413 Payload Too Large',{})
            payload=json.loads(env['wsgi.input'].read(length))
            return reply('200 OK',map_user_assets(payload) if path.endswith('/map-user') else mutate(payload))
        if method=='GET' and (path in ('/','/admin','/users') or path.startswith('/a/')):return reply('200 OK',(ROOT.parent/'public/index.html').read_bytes(),'text/html; charset=utf-8')
        return reply('404 Not Found',{})
    except (ValueError,KeyError) as e:return reply('409 Conflict',{'error':str(e)})
    except PermissionError:return reply('409 Conflict',{'error':'엑셀이 열려 있거나 파일 권한이 없습니다. 파일을 닫고 재시도하세요.'})
    except Exception:
        import traceback;traceback.print_exc();return reply('500 Internal Server Error',{'error':'처리 실패. 서버 로그를 확인하세요.'})

if __name__=='__main__':
    if not admin_service.settings(__import__(__name__))['password_configured'] and len(os.environ.get('APP_PASSWORD',''))<12:raise SystemExit('APP_PASSWORD를 12자 이상으로 지정하세요.')
    if os.environ.get('BACKEND_GATEWAY_TOKEN') and len(os.environ['BACKEND_GATEWAY_TOKEN'])<24:raise SystemExit('게이트웨이 토큰은 24자 이상이어야 합니다.')
    def stop_server(signum, frame): raise SystemExit(0)
    signal.signal(signal.SIGTERM,stop_server)
    lock=FILE.with_suffix('.app.lock')
    try: fd=os.open(lock,os.O_CREAT|os.O_EXCL|os.O_WRONLY)
    except FileExistsError:raise SystemExit('실행 중인 앱이 있습니다. 비정상 종료였다면 실행 프로세스가 없는지 확인한 뒤 .app.lock 파일을 삭제하세요.')
    try:
        os.close(fd);initialize();print('Bumjin Assets: '+os.environ.get('PUBLIC_BASE_URL','http://127.0.0.1:'+os.environ.get('PORT','1880'))+' (로그인: admin)')
        make_server(os.environ.get('HOST','127.0.0.1'),int(os.environ.get('PORT','1880')),application).serve_forever()
    finally:lock.unlink(missing_ok=True)
