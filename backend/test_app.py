import io, json, os, shutil, tempfile, unittest
from pathlib import Path
import openpyxl
import app

class PilotTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.original=app.FILE
        app.FILE=Path(self.tmp.name)/'test.xlsx';shutil.copy2(self.original,app.FILE)
        app.initialize()
    def tearDown(self):app.FILE=self.original;self.tmp.cleanup()
    def test_roundtrip_and_conflict(self):
        a=app.snapshot();self.assertGreater(len(a['assets']),300)
        ids=[x['id'] for x in a['assets']];self.assertEqual(len(ids),len(set(ids)))
        payload={'division':'BJC','number':'TEST-ONLY-001','user':'테스트 사용자','model':'=HYPERLINK("bad")','date':'2025-01-01','status':'보관','revision':a['revision']}
        result=app.mutate(payload);b=app.snapshot();new=next(x for x in b['assets'] if x['id']==result['id'])
        self.assertEqual(new['date_label'],'제조일자');self.assertEqual(new['user'],'테스트 사용자')
        with self.assertRaises(ValueError):app.mutate(payload)
        app.mutate({**new,'user':'두번째 사용자','status':'사용','revision':b['revision']})
        c=app.snapshot();changed=next(x for x in c['assets'] if x['id']==result['id'])
        self.assertEqual(changed['user'],'두번째 사용자');self.assertEqual(len(c['history']),2)
        w=openpyxl.load_workbook(app.FILE);original=openpyxl.load_workbook(self.original)
        self.assertTrue(set(original.sheetnames).issubset(w.sheetnames))
        for name in original.sheetnames:
            self.assertEqual(str(original[name].merged_cells),str(w[name].merged_cells))
        s=w['5.BJC PC 현황'];self.assertEqual(s[f'G{new["row"]}'].data_type,'s')
        self.assertTrue(s[f'V{new["row"]}'].value.startswith('=IF'))
        self.assertIn(b'<svg',app.qr_svg('https://assets.example/a/'+result['id']))
        app.initialize();self.assertEqual([x['id'] for x in app.snapshot()['assets']], [x['id'] for x in c['assets']])
    def test_http_auth_and_csrf(self):
        os.environ['APP_PASSWORD']='test-password-123'
        statuses=[]
        app.application({'PATH_INFO':'/api/data'},lambda s,h:statuses.append(s))
        self.assertTrue(statuses[-1].startswith('401'))
        import base64
        auth='Basic '+base64.b64encode(b'admin:test-password-123').decode()
        env={'PATH_INFO':'/api/data','REQUEST_METHOD':'GET','HTTP_AUTHORIZATION':auth}
        out=app.application(env,lambda s,h:statuses.append(s));self.assertTrue(statuses[-1].startswith('200'));self.assertIn('assets',json.loads(out[0]))
        app.application({**env,'PATH_INFO':'/api/assets','REQUEST_METHOD':'POST'},lambda s,h:statuses.append(s))
        self.assertTrue(statuses[-1].startswith('403'))

if __name__=='__main__':unittest.main()
