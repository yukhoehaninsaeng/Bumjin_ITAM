"""Excel organization tree; scoped LDAP department writes, never user/DN creation."""
import datetime as dt
import json
import os
import uuid
import openpyxl
import admin_service

SHEET='앱_조직'
AUDIT='앱_조직변경이력'

def rows(w):
    if SHEET not in w:return []
    return [dict(zip(('id','name','parent_id'),[str(v or '') for v in r])) for r in w[SHEET].iter_rows(min_row=2,max_col=3,values_only=True) if r[0]]

def history(w):
    if AUDIT not in w:return []
    return [list(map(lambda v:str(v or ''),r)) for r in w[AUDIT].iter_rows(min_row=max(2,w[AUDIT].max_row-49),max_col=7,values_only=True)][::-1]

def path_for(items,oid):
    index={o['id']:o for o in items};seen=set();parts=[]
    while oid:
        if oid in seen or oid not in index:raise ValueError('조직의 상위 관계가 올바르지 않습니다.')
        seen.add(oid);o=index[oid];parts.append(o['name']);oid=o['parent_id']
    return '/'.join(reversed(parts))

def import_paths(app,w,paths):
    items=rows(w);by_path={path_for(items,o['id']):o['id'] for o in items}
    for path in paths:
        parts=[x.strip() for x in str(path or '').split('/')]
        if not path or any(not x or len(x)>80 for x in parts) or len(parts)>12:continue
        parent=''
        for i,name in enumerate(parts):
            key='/'.join(parts[:i+1])
            if key not in by_path:
                oid=str(uuid.uuid4());items.append({'id':oid,'name':name,'parent_id':parent});by_path[key]=oid
            parent=by_path[key]
    write_rows(app,w,items)

def write_rows(app,w,items):
    if SHEET not in w:w.create_sheet(SHEET)
    s=w[SHEET];s.delete_rows(1,s.max_row);s.append(['조직ID','조직명','상위조직ID'])
    for o in items:
        r=s.max_row+1
        for c,k in enumerate(('id','name','parent_id'),1):app.literal(s.cell(r,c),o[k])

ASSET_ORGS='앱_자산조직'
IMPORT_STATE='앱_조직가져오기'

def asset_links(w):
    if ASSET_ORGS not in w:return {}
    return {str(r[0]):str(r[1]) for r in w[ASSET_ORGS].iter_rows(min_row=2,max_col=2,values_only=True) if r[0] and r[1]}

def bootstrap_from_excel(app,w):
    """One-time snapshot of original divisions/departments; stable links survive edits."""
    if IMPORT_STATE in w and w[IMPORT_STATE]['A1'].value=='EXCEL_ORG_IMPORT_V1':return False
    if IMPORT_STATE in w or ASSET_ORGS in w:raise ValueError('조직 가져오기용 시트가 이미 사용 중입니다. 관리자에게 확인하세요.')
    all_assets=app.assets(w);paths=['BJC','BJE','BJM'];asset_paths={}
    for a in all_assets:
        division=a['division'];department=a['department'].strip()
        parts=[x.strip() for x in department.split('/') if x.strip()]
        if parts and parts[0]==division:parts=parts[1:]
        # Do not guess hierarchy from spaces or department names.
        if any(len(x)>80 for x in parts) or len(parts)>11:parts=[department[:80]]
        path='/'.join([division]+parts);paths.append(path);asset_paths[a['id']]=path
    import_paths(app,w,paths)
    items=rows(w);lookup={path_for(items,o['id']):o['id'] for o in items}
    s=w.create_sheet(ASSET_ORGS);s.append(['자산ID','조직ID'])
    for aid,path in asset_paths.items():
        r=s.max_row+1;app.literal(s.cell(r,1),aid);app.literal(s.cell(r,2),lookup[path])
    marker=w.create_sheet(IMPORT_STATE);marker['A1']='EXCEL_ORG_IMPORT_V1'
    marker['A2']=dt.datetime.now().astimezone().isoformat()
    marker['A3']='사업부/부서 최초 가져오기 완료. 이후 조직도 편집은 원본 대장과 LDAP를 자동 변경하지 않습니다.'
    return True

def journal(app,event):
    # Durable recovery trail for cross-system writes; never contains credentials.
    p=app.FILE.parent/'ldap_department_audit.jsonl'
    fd=os.open(p,os.O_WRONLY|os.O_CREAT|os.O_APPEND,0o600)
    with os.fdopen(fd,'a',encoding='utf-8') as f:
        f.write(json.dumps(event,ensure_ascii=False)+'\n');f.flush();os.fsync(f.fileno())

def append_audit(app,w,values):
    row=w[AUDIT].max_row+1
    for col,value in enumerate(values,1):app.literal(w[AUDIT].cell(row,col),value)

def selected_users(app,w,payload):
    ids=payload.get('user_ids')
    if not isinstance(ids,list) or not 1<=len(ids)<=50 or any(not isinstance(x,str) for x in ids) or len(ids)!=len(set(ids)):
        raise ValueError('LDAP 사용자 1~50명을 중복 없이 선택하세요.')
    users={str(r[3]):(n,list(r)) for n,r in enumerate(w['앱_LDAP사용자'].iter_rows(min_row=2,max_col=7,values_only=True),2) if r[0] and r[3]}
    if any(i not in users for i in ids):raise ValueError('LDAP 사용자 목록이 변경되었습니다. 다시 동기화하세요.')
    return [users[i] for i in ids]

def mutate(app,payload):
    with app.LOCK:
        old=app.digest()
        if payload.get('revision')!=old:raise ValueError('다른 변경이 있습니다. 새로고침 후 다시 실행하세요.')
        w=openpyxl.load_workbook(app.FILE);items=rows(w);operation=payload.get('operation');oid=str(payload.get('id',''))
        if operation in ('create','update','delete'):
            before={o['id']:path_for(items,o['id']) for o in items}
            if operation=='delete':
                if not any(o['id']==oid for o in items):raise ValueError('조직이 없습니다.')
                if any(o['parent_id']==oid for o in items):raise ValueError('하위 조직을 먼저 이동하거나 삭제하세요.')
                if oid in asset_links(w).values():raise ValueError('자산이 연결된 조직은 삭제할 수 없습니다. 이름 또는 상위 조직을 수정하세요.')
                if any(str(r[0] or '')==before[oid] for r in w['앱_LDAP사용자'].iter_rows(min_row=2,min_col=5,max_col=5,values_only=True)):
                    raise ValueError('사용자가 있는 조직입니다. LDAP 소속을 먼저 변경하세요.')
                items=[o for o in items if o['id']!=oid]
            else:
                name=str(payload.get('name','')).strip();parent=str(payload.get('parent_id',''))
                if not name or len(name)>80 or '/' in name or any(ord(c)<32 for c in name):raise ValueError('조직명은 / 없이 1~80자로 입력하세요.')
                if parent and not any(o['id']==parent for o in items):raise ValueError('상위 조직이 없습니다.')
                if operation=='create':oid=str(uuid.uuid4());items.append({'id':oid,'name':name,'parent_id':parent})
                else:
                    target=next((o for o in items if o['id']==oid),None)
                    if not target:raise ValueError('조직이 없습니다.')
                    target.update(name=name,parent_id=parent)
                paths=[path_for(items,o['id']) for o in items]
                if len(set(paths))!=len(paths):raise ValueError('같은 상위 조직에 동일한 조직명이 있습니다.')
                if any(len(p)>500 or len(p.split('/'))>12 for p in paths):raise ValueError('조직 경로는 12단계, 500자 이내입니다.')
                changed={p for i,p in before.items() if path_for(items,i)!=p}
                if any(str(r[0] or '') in changed for r in w['앱_LDAP사용자'].iter_rows(min_row=2,min_col=5,max_col=5,values_only=True)):
                    raise ValueError('소속 사용자가 있는 조직은 바로 이름/상위를 변경할 수 없습니다. 새 조직을 만든 뒤 사용자를 배정하세요.')
            write_rows(app,w,items);app.save(w,old);return {'ok':True,'id':oid}
        if operation=='inspect':
            users=selected_users(app,w,payload)
            if len(users)!=1:raise ValueError('속성 확인은 LDAP 사용자 한 명만 선택하세요.')
            import ldap3
            from ldap3.utils.conv import escape_filter_chars
            cfg=admin_service.settings(app)['ldap'];conn=admin_service.connect_directory(app)
            try:
                stable=str(users[0][1][3]);filt='(&'+cfg['user_filter']+'('+cfg['id_attr']+'='+escape_filter_chars(stable)+'))'
                conn.search(cfg['user_base'] or cfg['base_dn'],filt,attributes=ldap3.ALL_ATTRIBUTES)
                entries=[e for e in conn.response if e['type']=='searchResEntry']
                if conn.result.get('result')!=0 or len(entries)!=1:raise ValueError('LDAP 사용자 한 명을 조회하지 못했습니다.')
                allowed={'uid','cn','description','displayname','mail','departmentnumber','department','ou','employeenumber'}
                return {'attributes':{k:v for k,v in entries[0]['attributes'].items() if k.lower() in allowed}}
            except ValueError:raise
            except Exception as e:raise ValueError(admin_service.ldap_error(e)) from None
            finally:conn.unbind()
        if operation not in ('preview','assign'):raise ValueError('조직 작업을 선택하세요.')
        path=path_for(items,oid) if oid else ''
        users=selected_users(app,w,payload)
        preview=[{'directory_id':str(r[3]),'uid':str(r[0]),'description':str(r[6] or ''),'before':str(r[4] or ''),'after':path} for _,r in users]
        if operation=='preview':return {'users':preview,'department':path,'attribute':admin_service.settings(app)['ldap']['department_attr']}
        cfg=admin_service.settings(app)['ldap']
        attr=cfg.get('department_attr','')
        if not cfg.get('allow_department_write'):raise ValueError('관리자 설정에서 LDAP 부서 쓰기를 활성화하세요.')
        if attr.lower() not in ('departmentnumber','department','ou'):raise ValueError('LDAP 부서 속성을 확인하세요. 허용: departmentNumber, department, ou')
        import ldap3
        from ldap3.utils.conv import escape_filter_chars
        conn=admin_service.connect_directory(app);results=[];batch=str(uuid.uuid4())
        if AUDIT not in w:w.create_sheet(AUDIT).append(['시각','고유ID','계정','이전부서','변경부서','결과','배치ID'])
        try:
            # Complete identity/read preflight before modifying any entry.
            targets=[]
            for row,r in users:
                filt='(&'+cfg['user_filter']+'('+cfg['id_attr']+'='+escape_filter_chars(str(r[3]))+'))'
                conn.search(cfg['user_base'] or cfg['base_dn'],filt,attributes=[attr,cfg['uid_attr'],cfg['id_attr']])
                entries=[e for e in conn.response if e['type']=='searchResEntry']
                if conn.result.get('result')!=0 or len(entries)!=1:raise ValueError('LDAP 사용자를 고유 ID로 확인할 수 없습니다. 다시 동기화하세요.')
                entry=entries[0];a=entry['attributes'];val=a.get(attr,[]);values=val if isinstance(val,list) else [val] if val else []
                if len(values)>1:raise ValueError('부서 속성에 값이 여러 개 있습니다. 한 값으로 정리 후 동기화하세요.')
                current=str(values[0]) if values else ''
                uid=a.get(cfg['uid_attr']);uid=uid[0] if isinstance(uid,list) and uid else uid
                if str(uid or '')!=str(r[0]) or current!=str(r[4] or ''):raise ValueError('LDAP 계정/부서가 외부에서 변경되었습니다. 다시 동기화하세요.')
                targets.append((row,r,entry['dn'],current))
            for row,r,dn,current in targets:
                event={'time':dt.datetime.now().astimezone().isoformat(),'directory_id':str(r[3]),'uid':str(r[0]),'before':current,'after':path,'batch':batch}
                if current==path:results.append({**event,'status':'unchanged'});continue
                journal(app,{**event,'status':'pending'})
                try:
                    ok=conn.modify(dn,{attr:[(ldap3.MODIFY_REPLACE,[path] if path else [])]})
                    if not ok or conn.result.get('result')!=0:raise ValueError('LDAP 수정 거부: '+str(conn.result.get('description','unknown')))
                except Exception as e:
                    status='failed' if isinstance(e,(ValueError,ldap3.core.exceptions.LDAPOperationResult)) else 'unknown'
                    message=str(e) if isinstance(e,ValueError) else admin_service.ldap_error(e)
                    if status=='unknown':message+=' 수정 응답을 확인하지 못했습니다. 다시 동기화해서 실제 부서를 확인하세요.'
                    journal(app,{**event,'status':status,'message':message})
                    results.append({**event,'status':status,'message':message})
                    append_audit(app,w,[event['time'],event['directory_id'],event['uid'],current,path,'실패' if status=='failed' else '확인 필요',batch])
                    break  # Report partial success; never silently retry or roll back other users.
                journal(app,{**event,'status':'success'})
                app.literal(w['앱_LDAP사용자'].cell(row,5),path)
                append_audit(app,w,[event['time'],event['directory_id'],event['uid'],current,path,'성공',batch])
                results.append({**event,'status':'success'})
            if results:
                try:app.save(w,old)
                except Exception: return {'results':results,'unprocessed':len(users)-len(results),'warning':'LDAP 결과는 기록됐지만 Excel 저장에 실패했습니다. 재실행하지 말고 LDAP를 다시 동기화하세요. data/ldap_department_audit.jsonl에서 결과를 확인하세요.'}
            return {'results':results,'unprocessed':len(users)-len(results)}
        except ValueError:raise
        except Exception as e:raise ValueError(admin_service.ldap_error(e)) from None
        finally:conn.unbind()
